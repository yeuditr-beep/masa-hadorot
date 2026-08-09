from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

p = Path(__file__).parent / "מסע-הדורות-מסמך-עריכה-לבונה-האתר.docx"
d = Document(p)
empty = sum(1 for t in d.tables if any(not c.text.strip() for row in t.rows for c in row.cells))
rtl = sum(1 for p in d.paragraphs if p._p.pPr is not None and p._p.pPr.find(qn("w:bidi")) is not None)
headings = sum(1 for p in d.paragraphs if p.style.name.startswith("Heading"))
print({"paragraphs": len(d.paragraphs), "tables": len(d.tables), "sections": len(d.sections), "empty_tables": empty, "rtl_paragraphs": rtl, "headings": headings, "bytes": p.stat().st_size})
